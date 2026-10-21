"""Step 3. Build the Excel model with live formulas from the simulated raw export.
Recalculate afterwards:  soffice --headless --convert-to xlsx --outdir model/recalc model/pipeline_presence_model.xlsx"""
import json, os, re, sys
import numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as CL

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RAW = os.path.join(ROOT, "data", "raw")
PY = json.load(open(os.path.join(ROOT, "outputs", "audit_results.json")))
OUT = os.path.join(ROOT, "model", "pipeline_presence_model.xlsx")
rd = lambda n: pd.read_csv(os.path.join(RAW, f"{n}.csv"), keep_default_na=False)
owners, acc, op, act, fb, ev, tg = (rd(n) for n in ("owners", "accounts", "opportunities", "activities", "feedback", "events", "targets"))
accp = pd.read_csv(os.path.join(ROOT, "data", "processed", "accounts_clean.csv"), keep_default_na=False)
fbp = pd.read_csv(os.path.join(ROOT, "data", "processed", "feedback_tagged.csv"), keep_default_na=False)

F = "Arial"
f_n, f_b = Font(name=F, size=10), Font(name=F, size=10, bold=True)
f_in, f_lk = Font(name=F, size=10, color="0000FF"), Font(name=F, size=10, color="008000")
f_h, f_t, f_note = Font(name=F, size=10, bold=True, color="FFFFFF"), Font(name=F, size=14, bold=True, color="12202B"), Font(name=F, size=9, italic=True, color="555555")
fill_h, fill_s, fill_k = PatternFill("solid", fgColor="12202B"), PatternFill("solid", fgColor="E4EEF0"), PatternFill("solid", fgColor="FFF2CC")
FM = dict(pct="0.0%", num="#,##0", n1="#,##0.0", n2="#,##0.00", x='0.00"x"', date="yyyy-mm-dd", int="0")
wb = Workbook(); wb.remove(wb.active)
def sh(name, tab="12202B"):
    w = wb.create_sheet(name); w.sheet_properties.tabColor = tab; w.sheet_view.showGridLines = False; return w
def put(w, r, c, v, font=None, fmt=None, fill=None, al=None):
    x = w.cell(r, c, v); x.font = font or f_n
    if fmt: x.number_format = FM.get(fmt, fmt)
    if fill: x.fill = fill
    if al: x.alignment = al
    return x
def hdr(w, r, labels, c0=1, h=40):
    for i, l in enumerate(labels):
        put(w, r, c0 + i, l, f_h, fill=fill_h, al=Alignment(horizontal="center", vertical="center", wrap_text=True))
    w.row_dimensions[r].height = h
def title(w, t, sub=None):
    put(w, 1, 1, t, f_t)
    if sub: put(w, 2, 1, sub, f_note)
def widths(w, d):
    for k, v in d.items(): w.column_dimensions[k].width = v
def dt(v): return None if v in ("", None) or pd.isna(v) else pd.Timestamp(v).to_pydatetime()

sheets = {n: sh(n, c) for n, c in [("README", "7A7A7A"), ("Dashboard_LoS", "E4572E"), ("Dashboard_Industry", "E4572E"), ("Priority_Accounts", "E4572E"), ("DQ_Scorecard", "E4572E"), ("Owner_Hygiene", "E4572E"),
                                    ("Feedback", "2E86AB"), ("Events", "2E86AB"), ("Inputs", "2F55D4"), ("Opps", "2F55D4"), ("Accounts", "2F55D4"), ("Activities", "2F55D4"), ("Owners", "2F55D4"), ("Followups", "7A7A7A"), ("Checks", "2E8B57")]}
# ------------------------------------------------------------------ Inputs
w = sheets["Inputs"]; title(w, "Inputs", "Blue = input. All dashboards read these cells. Data is SIMULATED (see README).")
put(w, 4, 1, "As-of date"); put(w, 4, 2, pd.Timestamp("2026-10-19").to_pydatetime(), f_in, "date")
put(w, 5, 1, "Stale threshold (days without activity)"); put(w, 5, 2, 60, f_in, "int")
put(w, 6, 1, "FY27 start"); put(w, 6, 2, pd.Timestamp("2026-07-01").to_pydatetime(), f_in, "date")
put(w, 7, 1, "FY27 end"); put(w, 7, 2, pd.Timestamp("2027-06-30").to_pydatetime(), f_in, "date")
put(w, 8, 1, "Look-back for win rate (days)"); put(w, 8, 2, 365, f_in, "int")
ASOF, STALE, FYS, FYE, LB = (f"Inputs!$B${r}" for r in (4, 5, 6, 7, 8))
hdr(w, 10, ["Stage", "Default probability %"])
for i, (s_, p_) in enumerate([("Prospecting", 10), ("Qualification", 25), ("Proposal", 50), ("Negotiation", 75), ("Closed Won", 100), ("Closed Lost", 0)]):
    put(w, 11 + i, 1, s_); put(w, 11 + i, 2, p_, f_in, "int")
hdr(w, 19, ["Line of service", "FY27 BD target (KES m) [simulated]"])
for i, r in tg.iterrows():
    put(w, 20 + i, 1, r.line_of_service); put(w, 20 + i, 2, float(r.fy27_bd_target_kes_m), f_in, "num")
w.cell(26, 1, "Targets are simulated planning inputs, not real firm targets.").font = f_note
widths(w, {"A": 44, "B": 30})
STG = "Inputs!$A$11:$A$16"; STP = "Inputs!$B$11:$B$16"

# ------------------------------------------------------------------ Owners
w = sheets["Owners"]; title(w, "Owners (simulated staff)")
hdr(w, 4, ["Owner ID", "Name", "Grade", "Home line of service", "Active?"])
for i, r in owners.iterrows():
    for c, v in enumerate([r.owner_id, r.owner_name, r.grade, r.los, "Yes" if str(r.active).lower() == "true" else "No"]):
        put(w, 5 + i, 1 + c, v, f_in)
OW0, OW1 = 5, 4 + len(owners)
widths(w, {"A": 10, "B": 22, "C": 16, "D": 20, "E": 10})

# ------------------------------------------------------------------ Accounts
w = sheets["Accounts"]; title(w, "Accounts", "Raw export + normalised-name key [E] from Python. Duplicate flag and master ID are live formulas.")
hdr(w, 4, ["Account ID", "Account name", "Industry (raw)", "Tier", "Type", "City", "Account driver", "Normalised name key [E]", "Duplicate?", "Master account ID", "Industry (clean)", "Missing industry?"])
accm = accp.set_index("account_id")
A0, A1 = 5, 4 + len(acc)
for i, r in acc.iterrows():
    rr = 5 + i
    for c, v in enumerate([r.account_id, r.account_name, r.industry, r.tier, r.account_type, r.city, r.account_driver]):
        put(w, rr, 1 + c, v if v != "" else None, f_in)
    key = re.sub(r"[^a-z0-9]", "", r.account_name.lower().replace(" limited", "").replace(" ltd", "").replace("county govt", "county government"))
    put(w, rr, 8, key, f_in)
    put(w, rr, 9, f"=IF(COUNTIF($H${A0}:$H${A1},H{rr})>1,1,0)", fmt="int")
    put(w, rr, 10, f"=INDEX($A${A0}:$A${A1},MATCH(H{rr},$H${A0}:$H${A1},0))")
    put(w, rr, 11, f'=IF(C{rr}<>"",C{rr},IF(INDEX($C${A0}:$C${A1},MATCH(J{rr},$A${A0}:$A${A1},0))<>"",INDEX($C${A0}:$C${A1},MATCH(J{rr},$A${A0}:$A${A1},0)),"Unassigned"))')
    put(w, rr, 12, f'=IF(C{rr}="",1,0)', fmt="int")
widths(w, {"A": 11, "B": 34, "C": 28, "D": 10, "E": 24, "F": 11, "G": 12, "H": 30, "I": 11, "J": 14, "K": 28, "L": 12})
w.freeze_panes = "C5"

# ------------------------------------------------------------------ Opps
w = sheets["Opps"]; title(w, "Opportunities: raw export plus live data-quality flags", "Columns A-N are the simulated Salesforce export. Columns O onward are formulas.")
H = ["Opp ID", "Account ID", "Opportunity", "Line of service", "Stage", "Amount (KES)", "Probability %", "Created", "Close date", "Last activity", "Owner ID", "Lead source", "Next step", "Lost reason",
     "Open?", "Industry", "Tier", "Stage probability %", "Weighted (KES)", "Days since activity", "Flag: no amount", "Flag: no owner", "Flag: owner left", "Flag: close date passed", "Flag: stale", "Flag: probability off",
     "Flag: no next step", "Flag: lost no reason", "Flags", "Clean?", "Closes in FY27 (open)", "Won FY27 YTD", "Closed in look-back", "Cycle days"]
hdr(w, 4, H, h=54)
O0, O1 = 5, 4 + len(op)
for i, r in op.iterrows():
    rr = 5 + i
    vals = [r.opp_id, r.account_id, r.opp_name, r.line_of_service, r.stage, None if r.amount_kes in ("", None) or pd.isna(pd.to_numeric(r.amount_kes, errors="coerce")) else float(r.amount_kes), float(r.probability),
            dt(r.created_date), dt(r.close_date), dt(r.last_activity_date), r.owner_id or None, r.lead_source, r.next_step or None, r.lost_reason or None]
    for c, v in enumerate(vals):
        put(w, rr, 1 + c, v, f_in, {6: "num", 7: "int", 8: "date", 9: "date", 10: "date"}.get(1 + c))
    put(w, rr, 15, f'=IF(OR(E{rr}="Closed Won",E{rr}="Closed Lost"),0,1)', fmt="int")
    put(w, rr, 16, f"=INDEX(Accounts!$K${A0}:$K${A1},MATCH(B{rr},Accounts!$A${A0}:$A${A1},0))")
    put(w, rr, 17, f"=INDEX(Accounts!$D${A0}:$D${A1},MATCH(B{rr},Accounts!$A${A0}:$A${A1},0))")
    put(w, rr, 18, f"=INDEX({STP},MATCH(E{rr},{STG},0))", fmt="int")
    put(w, rr, 19, f"=IF(F{rr}=\"\",0,F{rr})*R{rr}/100", fmt="num")
    put(w, rr, 20, f"={ASOF}-J{rr}", fmt="int")
    put(w, rr, 21, f'=IF(OR(F{rr}="",F{rr}=0),1,0)', fmt="int")
    put(w, rr, 22, f'=IF(K{rr}="",1,0)', fmt="int")
    put(w, rr, 23, f'=IF(AND(O{rr}=1,K{rr}<>"",COUNTIFS(Owners!$A${OW0}:$A${OW1},K{rr},Owners!$E${OW0}:$E${OW1},"No")>0),1,0)', fmt="int")
    put(w, rr, 24, f"=IF(AND(O{rr}=1,I{rr}<{ASOF}),1,0)", fmt="int")
    put(w, rr, 25, f"=IF(AND(O{rr}=1,T{rr}>{STALE}),1,0)", fmt="int")
    put(w, rr, 26, f"=IF(AND(O{rr}=1,G{rr}<>R{rr}),1,0)", fmt="int")
    put(w, rr, 27, f'=IF(AND(O{rr}=1,M{rr}=""),1,0)', fmt="int")
    put(w, rr, 28, f'=IF(AND(E{rr}="Closed Lost",N{rr}=""),1,0)', fmt="int")
    put(w, rr, 29, f"=SUM(U{rr}:AB{rr})", fmt="int")
    put(w, rr, 30, f"=IF(AC{rr}=0,1,0)", fmt="int")
    put(w, rr, 31, f"=IF(AND(O{rr}=1,I{rr}>={FYS},I{rr}<={FYE}),1,0)", fmt="int")
    put(w, rr, 32, f'=IF(AND(E{rr}="Closed Won",I{rr}>={FYS},I{rr}<={ASOF}),1,0)', fmt="int")
    put(w, rr, 33, f"=IF(AND(O{rr}=0,I{rr}>{ASOF}-{LB},I{rr}<={ASOF}),1,0)", fmt="int")
    put(w, rr, 34, f"=I{rr}-H{rr}", fmt="int")
w.freeze_panes = "D5"; w.auto_filter.ref = f"A4:AH{O1}"
widths(w, {"A": 9, "B": 10, "C": 38, "D": 14, "E": 13, "F": 13, "G": 10, "H": 11, "I": 11, "J": 11, "K": 9, "L": 18, "M": 18, "N": 18, "P": 28})
for c in range(15, 35):
    if CL(c) not in ("P",): w.column_dimensions[CL(c)].width = 11
w.column_dimensions["P"].width = 28
def OC(col): return f"Opps!${col}${O0}:${col}${O1}"

# ------------------------------------------------------------------ Activities
w = sheets["Activities"]; title(w, "Activities (touchpoints)")
hdr(w, 4, ["Activity ID", "Account ID", "Type", "Date", "Owner ID"])
for i, r in act.iterrows():
    for c, v in enumerate([r.activity_id, r.account_id, r.activity_type, dt(r.activity_date), r.owner_id]):
        put(w, 5 + i, 1 + c, v, f_in, "date" if c == 3 else None)
T0, T1 = 5, 4 + len(act)
widths(w, {"A": 11, "B": 11, "C": 28, "D": 12, "E": 9})

# ------------------------------------------------------------------ Dashboard_LoS
w = sheets["Dashboard_LoS"]; title(w, "Line of Service dashboard (KES millions)", "Reconstructs the monthly BD-meeting dashboard: pipeline, FY27 coverage, win rate, cycle time and data quality. Targets are simulated.")
hdr(w, 4, ["Line of service", "Open opps", "Open pipeline (KES m)", "Weighted pipeline (KES m)", "FY27-closing weighted (KES m)", "FY27 target (KES m)", "Won FY27 YTD (KES m)", "Remaining to target (KES m)", "Coverage (x)",
           "Closed, look-back", "Won, look-back", "Win rate", "Avg cycle, won (days)", "Open opps data-clean"], h=56)
LOS = list(tg.line_of_service)
for i, l in enumerate(LOS):
    r = 5 + i
    put(w, r, 1, f"=Inputs!A{20+i}", f_lk, None, None)
    put(w, r, 2, f"=COUNTIFS({OC('D')},A{r},{OC('O')},1)", fmt="int")
    put(w, r, 3, f"=SUMIFS({OC('F')},{OC('D')},A{r},{OC('O')},1)/1000000", fmt="n1")
    put(w, r, 4, f"=SUMIFS({OC('S')},{OC('D')},A{r},{OC('O')},1)/1000000", fmt="n1")
    put(w, r, 5, f"=SUMIFS({OC('S')},{OC('D')},A{r},{OC('AE')},1)/1000000", fmt="n1")
    put(w, r, 6, f"=Inputs!B{20+i}", f_lk, "num")
    put(w, r, 7, f"=SUMIFS({OC('F')},{OC('D')},A{r},{OC('AF')},1)/1000000", fmt="n1")
    put(w, r, 8, f"=MAX(F{r}-G{r},0)", fmt="n1")
    put(w, r, 9, f'=IF(H{r}=0,"",E{r}/H{r})', fmt="x", fill=fill_k)
    put(w, r, 10, f"=COUNTIFS({OC('D')},A{r},{OC('AG')},1)", fmt="int")
    put(w, r, 11, f'=COUNTIFS({OC("D")},A{r},{OC("AG")},1,{OC("E")},"Closed Won")', fmt="int")
    put(w, r, 12, f'=IF(J{r}=0,"",K{r}/J{r})', fmt="pct")
    put(w, r, 13, f'=AVERAGEIFS({OC("AH")},{OC("D")},A{r},{OC("E")},"Closed Won")', fmt="n1")
    put(w, r, 14, f"=COUNTIFS({OC('D')},A{r},{OC('O')},1,{OC('AD')},1)/B{r}", fmt="pct")
r = 5 + len(LOS)
put(w, r, 1, "Total", f_b)
for c in (2, 3, 4, 5, 6, 7, 8, 10, 11):
    put(w, r, c, f"=SUM({CL(c)}5:{CL(c)}{r-1})", f_b, "int" if c in (2, 10, 11) else "n1")
put(w, r, 9, f"=E{r}/H{r}", f_b, "x", fill=fill_k); put(w, r, 12, f"=K{r}/J{r}", f_b, "pct")
put(w, r, 13, f'=AVERAGEIFS({OC("AH")},{OC("E")},"Closed Won")', f_b, "n1"); put(w, r, 14, f"=COUNTIFS({OC('O')},1,{OC('AD')},1)/B{r}", f_b, "pct")
w.conditional_formatting.add("I5:I9", CellIsRule(operator="lessThan", formula=["1"], font=Font(name=F, size=10, bold=True, color="C00000")))
put(w, 12, 1, "Coverage = weighted pipeline expected to close in FY27 / (target - won year to date). Below 1.0x means the weighted pipeline alone cannot reach the target.", f_note)
ch = BarChart(); ch.type = "col"; ch.title = "Open vs weighted pipeline by line of service (KES m)"; ch.height, ch.width = 8.5, 17
ch.add_data(Reference(w, min_col=3, max_col=4, min_row=4, max_row=9), titles_from_data=True); ch.set_categories(Reference(w, min_col=1, min_row=5, max_row=9))
w.add_chart(ch, "A14")
ch2 = BarChart(); ch2.type = "col"; ch2.title = "FY27 coverage (x)"; ch2.height, ch2.width = 8.5, 12
ch2.add_data(Reference(w, min_col=9, min_row=4, max_row=9), titles_from_data=True); ch2.set_categories(Reference(w, min_col=1, min_row=5, max_row=9))
w.add_chart(ch2, "J14")
widths(w, {"A": 20}); [w.column_dimensions.__setitem__(CL(c), w.column_dimensions[CL(c)]) for c in range(2, 15)]
for c in range(2, 15): w.column_dimensions[CL(c)].width = 14

# ------------------------------------------------------------------ Dashboard_Industry
w = sheets["Dashboard_Industry"]; title(w, "Industry group dashboard (KES millions)", "Industry comes from the account record; accounts without an industry fall into 'Unassigned' until fixed.")
INDS = [x["industry"] for x in PY["by_industry"]]
hdr(w, 4, ["Industry", "Open opps", "Open pipeline (KES m)", "Weighted pipeline (KES m)", "Share of weighted", "Closed, look-back", "Won, look-back", "Win rate", "Open opps data-clean"], h=48)
for i, ind in enumerate(INDS):
    r = 5 + i
    put(w, r, 1, ind, f_in)
    put(w, r, 2, f"=COUNTIFS({OC('P')},A{r},{OC('O')},1)", fmt="int")
    put(w, r, 3, f"=SUMIFS({OC('F')},{OC('P')},A{r},{OC('O')},1)/1000000", fmt="n1")
    put(w, r, 4, f"=SUMIFS({OC('S')},{OC('P')},A{r},{OC('O')},1)/1000000", fmt="n1")
    put(w, r, 5, f"=D{r}/D${5+len(INDS)}", fmt="pct")
    put(w, r, 6, f"=COUNTIFS({OC('P')},A{r},{OC('AG')},1)", fmt="int")
    put(w, r, 7, f'=COUNTIFS({OC("P")},A{r},{OC("AG")},1,{OC("E")},"Closed Won")', fmt="int")
    put(w, r, 8, f'=IF(F{r}=0,"",G{r}/F{r})', fmt="pct")
    put(w, r, 9, f'=IF(B{r}=0,"",COUNTIFS({OC("P")},A{r},{OC("O")},1,{OC("AD")},1)/B{r})', fmt="pct")
r = 5 + len(INDS)
put(w, r, 1, "Total", f_b)
for c in (2, 3, 4, 6, 7): put(w, r, c, f"=SUM({CL(c)}5:{CL(c)}{r-1})", f_b, "int" if c in (2, 6, 7) else "n1")
put(w, r, 5, f"=SUM(E5:E{r-1})", f_b, "pct"); put(w, r, 8, f"=G{r}/F{r}", f_b, "pct")
w.conditional_formatting.add(f"D5:D{r-1}", DataBarRule(start_type="num", start_value=0, end_type="max", color="2E86AB"))
ch = BarChart(); ch.type = "bar"; ch.title = "Weighted pipeline by industry (KES m)"; ch.height, ch.width = 9, 18
ch.add_data(Reference(w, min_col=4, min_row=4, max_row=r - 1), titles_from_data=True); ch.set_categories(Reference(w, min_col=1, min_row=5, max_row=r - 1)); w.add_chart(ch, f"A{r+3}")
widths(w, {"A": 32}); 
for c in range(2, 10): w.column_dimensions[CL(c)].width = 15

# ------------------------------------------------------------------ Priority accounts
w = sheets["Priority_Accounts"]; title(w, "Priority accounts: pipeline and recency of contact", "Priority list is an input. Open opportunities and last touch are live formulas.")
hdr(w, 4, ["Account ID", "Account", "Industry", "Driver", "Open opps", "Weighted pipeline (KES m)", "Last touchpoint", "Days since touch", "Needs attention?"], h=44)
pri = accp[accp.tier == "Priority"].reset_index(drop=True)
P0, P1 = 5, 4 + len(pri)
for i, r_ in pri.iterrows():
    r = 5 + i
    put(w, r, 1, r_.account_id, f_in); put(w, r, 2, f"=INDEX(Accounts!$B${A0}:$B${A1},MATCH(A{r},Accounts!$A${A0}:$A${A1},0))"); put(w, r, 3, f"=INDEX(Accounts!$K${A0}:$K${A1},MATCH(A{r},Accounts!$A${A0}:$A${A1},0))")
    put(w, r, 4, f"=INDEX(Owners!$B${OW0}:$B${OW1},MATCH(INDEX(Accounts!$G${A0}:$G${A1},MATCH(A{r},Accounts!$A${A0}:$A${A1},0)),Owners!$A${OW0}:$A${OW1},0))")
    put(w, r, 5, f"=COUNTIFS({OC('B')},A{r},{OC('O')},1)", fmt="int")
    put(w, r, 6, f"=SUMIFS({OC('S')},{OC('B')},A{r},{OC('O')},1)/1000000", fmt="n2")
    put(w, r, 7, f'=IF(COUNTIFS(Activities!$B${T0}:$B${T1},A{r})=0,"",_xlfn.MAXIFS(Activities!$D${T0}:$D${T1},Activities!$B${T0}:$B${T1},A{r}))', fmt="date")
    put(w, r, 8, f'=IF(G{r}="","",{ASOF}-G{r})', fmt="int")
    put(w, r, 9, f'=IF(OR(E{r}=0,H{r}="",H{r}>90),"Yes","")', f_b)
r = P1 + 2
put(w, r, 1, "Priority accounts", f_b); put(w, r, 5, f"=COUNTA(A{P0}:A{P1})", fmt="int")
put(w, r + 1, 1, "With at least one open opportunity", f_b); put(w, r + 1, 5, f'=COUNTIF(E{P0}:E{P1},">0")', fmt="int")
put(w, r + 2, 1, "Touched in the last 90 days", f_b); put(w, r + 2, 5, f'=COUNTIFS(H{P0}:H{P1},"<=90")', fmt="int")
put(w, r + 3, 1, "Needing attention (no open opp, or no touch in 90 days)", f_b); put(w, r + 3, 5, f'=COUNTIF(I{P0}:I{P1},"Yes")', fmt="int")
PRI_SUM = r
w.conditional_formatting.add(f"I{P0}:I{P1}", CellIsRule(operator="equal", formula=['"Yes"'], fill=PatternFill("solid", bgColor="F8D7DA")))
widths(w, {"A": 12, "B": 34, "C": 28, "D": 22, "E": 11, "F": 16, "G": 14, "H": 12, "I": 14})

# ------------------------------------------------------------------ DQ scorecard
w = sheets["DQ_Scorecard"]; title(w, "Salesforce data-quality scorecard", "Eight opportunity rules and two account rules. Open-opportunity clean rate is the headline measure.")
hdr(w, 4, ["Rule", "Applies to", "Failing", "Population", "Failure rate", "What the owner should do"], h=36)
RULES = [("Amount missing or zero", "all opportunities", "U", f"=COUNTA({OC('A')})", "Add an estimated amount"), ("No owner assigned", "all opportunities", "V", f"=COUNTA({OC('A')})", "Assign an owner"),
         ("Owner has left or is inactive", "open", "W", f"=SUM({OC('O')})", "Reassign the opportunity"), ("Open with close date in the past", "open", "X", f"=SUM({OC('O')})", "Update the close date or close it"),
         ("No activity for more than the stale threshold", "open", "Y", f"=SUM({OC('O')})", "Log an activity or confirm status"), ("Probability does not match stage", "open", "Z", f"=SUM({OC('O')})", "Correct the probability"),
         ("No next step recorded", "open", "AA", f"=SUM({OC('O')})", "Record the next step"), ("Lost without a reason", "closed lost", "AB", f'=COUNTIFS({OC("E")},"Closed Lost")', "Add the lost reason")]
for i, (lab, scope, col, pop, act_) in enumerate(RULES):
    r = 5 + i
    put(w, r, 1, lab); put(w, r, 2, scope); put(w, r, 3, f"=SUM({OC(col)})", fmt="int"); put(w, r, 4, pop, fmt="int"); put(w, r, 5, f"=C{r}/D{r}", fmt="pct", fill=fill_k); put(w, r, 6, act_)
w.conditional_formatting.add("E5:E12", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=0.4, color="E4572E"))
put(w, 14, 1, "Account: missing industry"); put(w, 14, 3, f"=SUM(Accounts!$L${A0}:$L${A1})", fmt="int"); put(w, 14, 4, f"=COUNTA(Accounts!$A${A0}:$A${A1})", fmt="int"); put(w, 14, 5, "=C14/D14", fmt="pct")
put(w, 15, 1, "Account: part of a duplicate pair"); put(w, 15, 3, f"=SUM(Accounts!$I${A0}:$I${A1})", fmt="int"); put(w, 15, 4, f"=COUNTA(Accounts!$A${A0}:$A${A1})", fmt="int"); put(w, 15, 5, "=C15/D15", fmt="pct")
put(w, 17, 1, "Open opportunities with no data issue", f_b); put(w, 17, 3, f"=COUNTIFS({OC('O')},1,{OC('AD')},1)", f_b, "int"); put(w, 17, 4, f"=SUM({OC('O')})", f_b, "int"); put(w, 17, 5, "=C17/D17", f_b, "pct", fill=fill_k)
put(w, 18, 1, "Open opportunities with at least one issue"); put(w, 18, 3, "=D17-C17", fmt="int"); put(w, 18, 5, "=C18/D17", fmt="pct")
put(w, 19, 1, "Issue count to chase (all rules, all opportunities)"); put(w, 19, 3, f"=SUM({OC('AC')})", fmt="int")
ch = BarChart(); ch.type = "bar"; ch.title = "Failure rate by rule"; ch.height, ch.width = 9, 18
ch.add_data(Reference(w, min_col=5, min_row=4, max_row=12), titles_from_data=True); ch.set_categories(Reference(w, min_col=1, min_row=5, max_row=12)); w.add_chart(ch, "A22")
widths(w, {"A": 48, "B": 18, "C": 10, "D": 12, "E": 12, "F": 40})

# ------------------------------------------------------------------ Owner hygiene
w = sheets["Owner_Hygiene"]; title(w, "Owner hygiene: who has opportunities to fix", "Use this to prepare the follow-up calls and the Salesforce clinics. Open opportunities only.")
hdr(w, 4, ["Owner ID", "Name", "Grade", "Active?", "Open opps", "Opps with an issue", "Clean rate", "Issues to fix"], h=36)
for i in range(len(owners)):
    r = 5 + i
    put(w, r, 1, f"=Owners!A{OW0+i}", f_lk); put(w, r, 2, f"=Owners!B{OW0+i}", f_lk); put(w, r, 3, f"=Owners!C{OW0+i}", f_lk); put(w, r, 4, f"=Owners!E{OW0+i}", f_lk)
    put(w, r, 5, f"=COUNTIFS({OC('K')},A{r},{OC('O')},1)", fmt="int")
    put(w, r, 6, f"=COUNTIFS({OC('K')},A{r},{OC('O')},1,{OC('AD')},0)", fmt="int")
    put(w, r, 7, f'=IF(E{r}=0,"",1-F{r}/E{r})', fmt="pct")
    put(w, r, 8, f"=SUMIFS({OC('AC')},{OC('K')},A{r},{OC('O')},1)", fmt="int")
r = 5 + len(owners)
put(w, r, 2, "(no owner)", f_b); put(w, r, 5, f'=COUNTIFS({OC("K")},"",{OC("O")},1)', fmt="int"); put(w, r, 6, f'=COUNTIFS({OC("K")},"",{OC("O")},1,{OC("AD")},0)', fmt="int"); put(w, r, 8, f'=SUMIFS({OC("AC")},{OC("K")},"",{OC("O")},1)', fmt="int")
put(w, r + 1, 2, "Total", f_b)
for c in (5, 6, 8): put(w, r + 1, c, f"=SUM({CL(c)}5:{CL(c)}{r})", f_b, "int")
OH_TOT = r + 1
w.conditional_formatting.add(f"G5:G{r-1}", CellIsRule(operator="lessThan", formula=["0.4"], font=Font(name=F, size=10, bold=True, color="C00000")))
widths(w, {"A": 10, "B": 24, "C": 16, "D": 9, "E": 11, "F": 14, "G": 11, "H": 12})

# ------------------------------------------------------------------ Feedback
w = sheets["Feedback"]; title(w, "Client feedback consolidation (simulated survey)", "Theme tags come from keyword matching in Python [E]. Summary tables are live formulas.")
hdr(w, 4, ["Feedback ID", "Account ID", "Line of service", "Survey date", "Rating (1-5)", "Comment", "Follow-up status", "Theme [E]", "Sentiment"], h=36)
for i, r_ in fbp.iterrows():
    r = 5 + i
    for c, v in enumerate([r_.feedback_id, r_.account_id, r_.line_of_service, dt(r_.survey_date), float(r_.rating), r_.comment, r_.follow_up_status, r_.theme]):
        put(w, r, 1 + c, v, f_in, "date" if c == 3 else None)
    put(w, r, 9, f'=IF(E{r}>=4,"Positive",IF(E{r}<=2,"Negative","Neutral"))')
B0, B1 = 5, 4 + len(fbp)
hdr(w, 4, ["Line of service", "Responses", "Average rating", "% positive", "% negative"], c0=12, h=36)
for i, l in enumerate(LOS):
    r = 5 + i
    put(w, r, 12, l); put(w, r, 13, f"=COUNTIFS($C${B0}:$C${B1},L{r})", fmt="int"); put(w, r, 14, f"=AVERAGEIFS($E${B0}:$E${B1},$C${B0}:$C${B1},L{r})", fmt="n2")
    put(w, r, 15, f'=COUNTIFS($C${B0}:$C${B1},L{r},$I${B0}:$I${B1},"Positive")/M{r}', fmt="pct"); put(w, r, 16, f'=COUNTIFS($C${B0}:$C${B1},L{r},$I${B0}:$I${B1},"Negative")/M{r}', fmt="pct")
put(w, 11, 12, "All", f_b); put(w, 11, 13, f"=COUNTA($A${B0}:$A${B1})", f_b, "int"); put(w, 11, 14, f"=AVERAGE($E${B0}:$E${B1})", f_b, "n2")
put(w, 11, 15, f'=COUNTIF($I${B0}:$I${B1},"Positive")/M11', f_b, "pct"); put(w, 11, 16, f'=COUNTIF($I${B0}:$I${B1},"Negative")/M11', f_b, "pct")
put(w, 12, 12, "Net sentiment (points, % positive - % negative)", f_b); put(w, 12, 16, "=(O11-P11)*100", f_b, "n1", fill=fill_k)
put(w, 13, 12, "Open follow-ups"); put(w, 13, 16, f'=COUNTIF($G${B0}:$G${B1},"Open")', fmt="int")
hdr(w, 16, ["Theme", "Positive", "Neutral", "Negative", "Total"], c0=12, h=30)
THEMES = ["responsiveness", "insight", "value", "communication", "digital"]
for i, t in enumerate(THEMES):
    r = 17 + i
    put(w, r, 12, t)
    for j, s_ in enumerate(["Positive", "Neutral", "Negative"]): put(w, r, 13 + j, f'=COUNTIFS($H${B0}:$H${B1},L{r},$I${B0}:$I${B1},"{s_}")', fmt="int")
    put(w, r, 16, f"=SUM(M{r}:O{r})", fmt="int")
widths(w, {"A": 10, "B": 10, "C": 14, "D": 11, "E": 10, "F": 56, "G": 12, "H": 16, "I": 10, "L": 36, "M": 11, "N": 13, "O": 11, "P": 12})

# ------------------------------------------------------------------ Events
w = sheets["Events"]; title(w, "Event support: funnel by event (simulated)", "Invited to registered to attended to opportunities created.")
hdr(w, 4, ["Event", "Industry", "Date", "Invited", "Registered", "Attended", "Opps created", "Registration rate", "Attendance rate", "Opps per 10 attendees"], h=44)
for i, r_ in ev.iterrows():
    r = 5 + i
    for c, v in enumerate([r_.event_name, r_.industry, dt(r_.date), int(r_.invited), int(r_.registered), int(r_.attended), int(r_.opps_created)]):
        put(w, r, 1 + c, v, f_in, "date" if c == 2 else None)
    put(w, r, 8, f"=IF(D{r}=0,\"\",E{r}/D{r})", fmt="pct"); put(w, r, 9, f'=IF(E{r}=0,"",F{r}/E{r})', fmt="pct"); put(w, r, 10, f'=IF(F{r}=0,"",G{r}/F{r}*10)', fmt="n1")
r = 5 + len(ev)
put(w, r, 1, "Held events (excl. planned)", f_b)
for c in (4, 5, 6, 7): put(w, r, c, f"=SUM({CL(c)}5:{CL(c)}{r-2})", f_b, "int")
put(w, r, 8, f"=E{r}/D{r}", f_b, "pct"); put(w, r, 9, f"=F{r}/E{r}", f_b, "pct"); put(w, r, 10, f"=G{r}/F{r}*10", f_b, "n1")
put(w, r + 2, 1, "The last row is the planned annual conference: see comms/event_pack.md for the run-of-show, checklist and success measures.", f_note)
widths(w, {"A": 40, "B": 28, "C": 12, "D": 10, "E": 11, "F": 10, "G": 12, "H": 14, "I": 14, "J": 16})

# ------------------------------------------------------------------ Followups (static from Python)
w = sheets["Followups"]; title(w, "Owner follow-up list (static output of analysis/audit.py)", "One row per issue per opportunity. Re-run the audit to refresh. High = past close date, owner left, no owner, or priority account.")
fu = pd.read_csv(os.path.join(ROOT, "outputs", "owner_followups.csv"), keep_default_na=False)
hdr(w, 4, list(fu.columns))
for i, r_ in fu.iterrows():
    for c, v in enumerate(r_.tolist()): put(w, 5 + i, 1 + c, v if v != "" else None, f_n, "num" if fu.columns[c] == "amount_kes" else None)
w.auto_filter.ref = f"A4:H{4+len(fu)}"
widths(w, {"A": 22, "B": 10, "C": 44, "D": 14, "E": 14, "F": 38, "G": 44, "H": 10})

# ------------------------------------------------------------------ Checks
w = sheets["Checks"]; title(w, "Checks", "Workbook formulas against the independent Python audit (analysis/audit.py). Every row should be TRUE.")
hdr(w, 4, ["Check", "Workbook", "Python", "Difference", "Pass"])
CK = []
def chk(label, wbv, pyv, tol=1e-6):
    r = 5 + len(CK); put(w, r, 1, label); put(w, r, 2, wbv, fmt="0.0000"); put(w, r, 3, pyv, f_in, "0.0000"); put(w, r, 4, f"=ABS(B{r}-C{r})", fmt="0.000000"); put(w, r, 5, f"=D{r}<={tol}", f_b); CK.append(r)
chk("Opportunities", f"=COUNTA({OC('A')})", PY["n_opps"]); chk("Open opportunities", f"=SUM({OC('O')})", PY["n_open"])
chk("Open pipeline (KES m)", "=Dashboard_LoS!C10", PY["open"]["open_kes_m"], 1e-6); chk("Weighted pipeline (KES m)", "=Dashboard_LoS!D10", PY["open"]["weighted_kes_m"], 1e-6)
chk("FY27 won year to date (KES m)", "=Dashboard_LoS!G10", PY["fy27"]["won_ytd_kes_m"]); chk("FY27 weighted closing pipeline (KES m)", "=Dashboard_LoS!E10", PY["fy27"]["pipeline_weighted_kes_m"])
chk("Win rate, look-back", "=Dashboard_LoS!L10", PY["win_rate_12m"])
for k, d in enumerate(PY["dq"]): chk(f"DQ failing: {d['label']}", f"=DQ_Scorecard!C{5+k}", d["failing"])
chk("Open opportunities with no issue (rate)", "=DQ_Scorecard!E17", PY["open_clean_rate"]); chk("Accounts missing industry", "=DQ_Scorecard!C14", PY["acc_dq"][0]["failing"]); chk("Accounts in duplicate pairs", "=DQ_Scorecard!C15", PY["acc_dq"][1]["failing"])
for k, b in enumerate(PY["by_los"]):
    chk(f"{b['los']}: weighted pipeline", f"=Dashboard_LoS!D{5+k}", b["weighted_kes_m"]); chk(f"{b['los']}: coverage (x)", f"=Dashboard_LoS!I{5+k}", b["coverage_x"]); chk(f"{b['los']}: win rate", f"=Dashboard_LoS!L{5+k}", b["win_rate"])
    chk(f"{b['los']}: avg cycle (days)", f"=Dashboard_LoS!M{5+k}", b["avg_cycle_days_won"]); chk(f"{b['los']}: open opps clean rate", f"=Dashboard_LoS!N{5+k}", b["dq_clean_rate"])
for k, b in enumerate(PY["by_industry"]): chk(f"Industry {b['industry']}: weighted pipeline", f"=Dashboard_Industry!D{5+k}", b["weighted_kes_m"])
chk("Priority accounts with an open opportunity", f"=Priority_Accounts!E{PRI_SUM+1}", PY["priority"]["with_open_opp"]); chk("Priority accounts touched in 90 days", f"=Priority_Accounts!E{PRI_SUM+2}", PY["priority"]["touched_90d"])
chk("Owner hygiene: issues across owners equals open-opportunity issues", f"=Owner_Hygiene!H{OH_TOT}", float(sum(1 for _ in range(0))) + float(pd.read_csv(os.path.join(RAW, '..', 'processed', 'opportunities_clean.csv'), keep_default_na=False).query("stage not in ['Closed Won','Closed Lost']").n_flags.sum()))
chk("Feedback: average rating", "=Feedback!N11", PY["feedback"]["avg"]); chk("Feedback: net sentiment points", "=Feedback!P12", PY["feedback"]["nps_like"], 1e-6)
chk("Industry total equals LoS total (open pipeline)", "=Dashboard_Industry!C" + str(5 + len(INDS)), "=Dashboard_LoS!C10")
w.cell(5 + len(CK) + 1, 1, "OVERALL").font = f_b
put(w, 5 + len(CK) + 1, 5, f'=IF(COUNTIF(E5:E{4+len(CK)},TRUE)={len(CK)},"ALL CHECKS PASS","CHECK FAILURES")', f_b, fill=fill_k)
widths(w, {"A": 66, "B": 14, "C": 14, "D": 14, "E": 20})

# ------------------------------------------------------------------ README
w = sheets["README"]; title(w, "Pipeline and Presence: CRM and marketing-intelligence model", "SIMULATED data for a fictional professional-services firm in Nairobi. No real client, firm or staff data is used.")
lines = [("Purpose", "Reconstructs the dashboards, data-quality checks and follow-ups that a Clients and Markets support associate would run from Salesforce, with every number a live formula."),
         ("Read first", "Dashboard_LoS and DQ_Scorecard. Then Priority_Accounts, Dashboard_Industry, Owner_Hygiene, Feedback and Events."),
         ("Inputs", "Inputs sheet: as-of date, stale threshold, FY27 dates, stage probabilities and simulated targets. Change them and every sheet updates."),
         ("Data", "Opps, Accounts, Activities and Owners are the simulated Salesforce export. Opps columns O onward are formula flags."),
         ("Followups", "Static list generated by analysis/audit.py for owner chasing."),
         ("Checks", "The workbook is reconciled to an independent Python audit."), ("Colour code", "Blue = input or raw data. Black = formula. Green = link to another sheet. Yellow fill = key output."),
         ("Tags", "[E] = computed in Python (normalised name key, feedback theme tags).")]
for i, (a, b) in enumerate(lines):
    put(w, 4 + i, 1, a, f_b); put(w, 4 + i, 2, b, al=Alignment(wrap_text=True, vertical="top"))
widths(w, {"A": 16, "B": 120})
wb.move_sheet("Dashboard_LoS", offset=0)
wb.save(OUT); print("saved", OUT, wb.sheetnames)
